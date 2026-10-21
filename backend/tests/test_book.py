"""Whole-book translation against a fake TranslateGemma server."""
import json
import re
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest

pytest.importorskip("bs4")

from iak import api, book, servers, translate  # noqa: E402
from iak.library import Library  # noqa: E402

from .test_api import TOKEN, call, wait  # noqa: E402


def shout(text: str) -> str:
    """The fake "translation": capitals outside tags and entities."""
    return "".join(p if p.startswith(("<", "&")) else p.upper()
                   for p in re.split(r"(<[^>]+>|&[#\w]+;)", text))


class FakeGemma:
    """/completion for TranslateGemma's prompt, the explicit one and revisions.
    keep_groups=False drops the <seg> tags of grouped units."""

    def __init__(self, keep_groups=True, n_ctx=8192, delay=0.0):
        self.prompts = []
        self.keep_groups = keep_groups
        outer = self

        class H(BaseHTTPRequestHandler):
            def _send(self, obj):
                body = json.dumps(obj).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(body)

            def do_GET(self):
                if self.path == "/props":
                    self._send({"chat_template": "{{ source_lang_code }}",
                                "default_generation_settings": {"n_ctx": n_ctx}})
                else:
                    self._send({"data": [{"id": "translategemma"}]})

            def do_POST(self):
                req = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                p = req["prompt"].split("<end_of_turn>")[0]
                outer.prompts.append(p)
                time.sleep(delay)
                if "REVISED" in p:
                    out = p.split(" TRANSLATION:\n", 1)[1].rsplit("\n\nREVISED", 1)[0]
                    out = out.replace("É", "E")                # the revision "corrects"
                elif " TEXT:\n" in p:
                    out = shout(p.split(" TEXT:\n", 1)[1])
                else:
                    out = shout(p.split(":\n\n\n", 1)[1])
                if not outer.keep_groups and out.count("<seg") > 1:
                    out = re.sub(r"</?seg[^>]*>", "", out)
                self._send({"content": out, "stop_type": "eos"})

            def log_message(self, *a):
                pass

        self.srv = HTTPServer(("127.0.0.1", 0), H)
        self.url = f"http://127.0.0.1:{self.srv.server_port}"
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def translator(self):
        return translate.Translator.connect(servers.Client(self.url))

    def close(self):
        self.srv.shutdown()


@pytest.fixture
def gemma():
    g = FakeGemma()
    yield g
    g.close()


CHAPTER = """<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="fr" xml:lang="fr">
<head><title>Chapitre {n}</title></head>
<body>
<h1 id="c{n}">Chapitre {n}</h1>
<p>Le registre de la bibliothèque mentionne l'achat d'un <em>atlas</em> et de cartes&nbsp;marines<a epub:type="noteref" href="#n{n}">1</a>.</p>
<p><span id="page{n}"/>Le relieur de la rue Saint-Jacques l'a réparé au printemps suivant.</p>
<p>Une ligne<br/>coupée en deux.</p>
<aside epub:type="footnote" id="n{n}"><p>Note du traducteur numéro {n}.</p></aside>
<pre>code non traduit</pre>
</body></html>
"""


def make_epub(path: Path, chapters=2) -> Path:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>""")
        items = "".join(f'<item id="c{i}" href="c{i}.xhtml" media-type="application/xhtml+xml"/>'
                        for i in range(1, chapters + 1))
        refs = "".join(f'<itemref idref="c{i}"/>' for i in range(1, chapters + 1))
        z.writestr("OEBPS/content.opf", f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Registre</dc:title>
<dc:language>fr</dc:language><dc:identifier id="id">x</dc:identifier></metadata>
<manifest>{items}<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest>
<spine toc="ncx">{refs}</spine></package>""")
        z.writestr("OEBPS/toc.ncx", """<?xml version="1.0" encoding="utf-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><navMap>
<navPoint id="p1"><navLabel><text>Chapitre un</text></navLabel><content src="c1.xhtml"/></navPoint>
</navMap></ncx>""")
        for i in range(1, chapters + 1):
            z.writestr(f"OEBPS/c{i}.xhtml", CHAPTER.format(n=i))
    return path


def settings_for(path: Path, tgt="en", **kw) -> dict:
    return {"name": path.name, "kind": book.FORMATS[path.suffix.lower()],
            "source": str(path), "sha256": book.file_sha256(path), "page_texts": {},
            "pages": 0, "doc_id": None, "src": "fr", "tgt": tgt, "passes": 2,
            "glossary": "", "limit": 0, **kw}


def test_epub_layout_kept(tmp_path, gemma):
    src = make_epub(tmp_path / "registre.epub")
    job = book.prepare(tmp_path / "lib", settings_for(src))
    assert book.run_job(job, gemma.translator(), lambda: False) == "done"
    out = job.output
    assert out.name == "registre.en.epub" and job.data["pct"] == 100
    with zipfile.ZipFile(out) as z:
        first = z.infolist()[0]
        assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
        c1 = z.read("OEBPS/c1.xhtml").decode()
        opf = z.read("OEBPS/content.opf").decode()
        ncx = z.read("OEBPS/toc.ncx").decode()
    assert "<dc:language>en</dc:language>" in opf and 'xml:lang="en"' in c1
    assert "<em>ATLAS</em>" in c1                           # inline formatting kept
    assert 'href="#n1">1</a>' in c1 and '<span id="page1"/>' in c1
    assert " " in c1 or "&#160;" in c1                 # the no-break space
    assert "UNE LIGNE<br/>COUPEE EN DEUX." in c1            # revision pass applied (É -> E)
    assert "NOTE DU TRADUCTEUR NUMERO 1." in c1
    assert "<pre>code non traduit</pre>" in c1
    assert "CHAPITRE UN" in ncx and "<title>CHAPITRE 1</title>" in c1
    # grouped paragraphs: few calls, TranslateGemma's own prompt for pass 1
    assert any(p.startswith("<start_of_turn>user\nYou are a professional French (fr)")
               for p in gemma.prompts)
    assert job.data["stats"]["group_fallbacks"] == 0


def test_pause_and_resume(tmp_path):
    g = FakeGemma(n_ctx=2048)
    try:
        src = make_epub(tmp_path / "long.epub", chapters=6)
        job = book.prepare(tmp_path / "lib", settings_for(src, passes=1))
        state = {"n": 0}

        def pause_after_two():
            state["n"] += 1
            return state["n"] > 2
        assert book.run_job(job, g.translator(), pause_after_two) == "paused"
        assert job.data["done_units"] == 2 and job.partial.exists()
        with zipfile.ZipFile(job.partial) as z:
            c1, c6 = z.read("OEBPS/c1.xhtml").decode(), z.read("OEBPS/c6.xhtml").decode()
        assert "ATLAS" in c1 and "atlas" in c6               # translated, then original
        calls = len(g.prompts)
        job = book.Job(job.folder)                          # as after a restart
        assert book.run_job(job, g.translator(), lambda: False) == "done"
        assert job.data["stats"]["cached"] >= 2 and not job.partial
        assert len(g.prompts) > calls
        with zipfile.ZipFile(job.output) as z:
            assert "ATLAS" in z.read("OEBPS/c6.xhtml").decode()
        assert "Paused after 2/" in "\n".join(job.log_tail(100))
    finally:
        g.close()


def test_group_fallback(tmp_path):
    g = FakeGemma(keep_groups=False)
    try:
        src = make_epub(tmp_path / "r.epub", chapters=3)
        job = book.prepare(tmp_path / "lib", settings_for(src, passes=1))
        assert book.run_job(job, g.translator(), lambda: False) == "done"
        s = job.data["stats"]
        assert s["group_fallbacks"] == 3                   # then paragraph by paragraph
        with zipfile.ZipFile(job.output) as z:
            assert "<em>ATLAS</em>" in z.read("OEBPS/c3.xhtml").decode()
        assert "paragraph by paragraph from now on" in "\n".join(job.log_tail())
    finally:
        g.close()


def test_pdf_with_ocr_pages(tmp_path, gemma):
    pymupdf = pytest.importorskip("pymupdf")
    path = tmp_path / "livre.pdf"
    doc = pymupdf.open()
    words = ["atlas", "relieur", "cartes", "registre", "volumes"]
    for i in range(5):
        page = doc.new_page(width=595, height=842)
        page.insert_text((60, 40), "Histoire du registre", fontsize=9)    # running head
        if i == 0:
            page.insert_text((60, 120), "Premier chapitre", fontsize=24)
        page.insert_textbox(pymupdf.Rect(60, 150, 535, 700),
                            f"Texte natif sur les {words[i]} de cette page, avec une "
                            f"longue phrase qui continue sur la page suivante sans", fontsize=12)
        page.insert_textbox(pymupdf.Rect(60, 700, 535, 780),
                            f"rupture, comme dans un vrai livre {words[i]}s.", fontsize=12)
        page.insert_text((300, 820), str(i + 1), fontsize=9)               # page number
    doc.save(path)
    # page 3 read by OCR in the library
    job = book.prepare(tmp_path / "lib", settings_for(
        path, page_texts={"3": "Texte lu par\nl'OCR de la page trois."}))
    assert book.run_job(job, gemma.translator(), lambda: False) == "done"
    # a PDF comes back as a PDF by default
    assert job.output.name == "livre.en.pdf" and job.data["out_format"] == "pdf"
    with pymupdf.open(job.output) as out:
        text = " ".join(p.get_text() for p in out)
        assert out.metadata["title"] == "livre" and out[0].rect.width < 450   # A5
    assert "PREMIER CHAPITRE" in text and "OCR DE LA PAGE TROIS" in text
    # the same translation in Markdown, from the cache: no model call
    calls = len(gemma.prompts)
    md_path = book.render(job, "md", tmp_path / "livre.en.md")
    assert len(gemma.prompts) == calls
    md = md_path.read_text(encoding="utf-8")
    assert md.startswith("# PREMIER CHAPITRE")
    assert "HISTOIRE DU REGISTRE" not in md and "\n\n3\n" not in md
    assert "TEXTE LU PAR L'OCR DE LA PAGE TROIS." in md
    assert ("TEXTE NATIF SUR LES ATLAS DE CETTE PAGE, AVEC UNE LONGUE PHRASE QUI CONTINUE "
            "SUR LA PAGE SUIVANTE SANS RUPTURE, COMME DANS UN VRAI LIVRE ATLASS.") in md
    assert md.count("TEXTE NATIF") == 4


def test_markdown_txt_docx(tmp_path, gemma):
    md = tmp_path / "notes.md"
    md.write_text("# Titre\n\nUn paragraphe\nsur deux lignes.\n\n- premier point\n"
                  "- second point\n\n```\ngarder ce code\n```\n\n> Une citation.\n",
                  encoding="utf-8")
    job = book.prepare(tmp_path / "lib", settings_for(md, passes=1))
    assert book.run_job(job, gemma.translator(), lambda: False) == "done"
    assert job.output.read_text(encoding="utf-8") == (
        "# TITRE\n\nUN PARAGRAPHE SUR DEUX LIGNES.\n\n- PREMIER POINT\n- SECOND POINT\n\n"
        "```\ngarder ce code\n```\n\n> UNE CITATION.\n")

    txt = tmp_path / "journal.txt"
    txt.write_text("Première ligne.\n\nDeuxième paragraphe.\n", encoding="utf-8")
    job = book.prepare(tmp_path / "lib", settings_for(txt, passes=1))
    assert book.run_job(job, gemma.translator(), lambda: False) == "done"
    assert job.output.read_text(encoding="utf-8") == "PREMIÈRE LIGNE.\n\nDEUXIÈME PARAGRAPHE.\n"

    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("Le registre", 1)
    p = d.add_paragraph("Un ")
    p.add_run("atlas").italic = True
    p.add_run(" endommagé.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cellule gauche"
    t.cell(0, 1).text = "Cellule droite"
    path = tmp_path / "notes.docx"
    d.save(path)
    job = book.prepare(tmp_path / "lib", settings_for(path, passes=1))
    assert book.run_job(job, gemma.translator(), lambda: False) == "done"
    out = docx.Document(str(job.output))
    texts = [p.text for p in out.paragraphs]
    assert texts[:2] == ["LE REGISTRE", "UN ATLAS ENDOMMAGÉ."]
    assert out.paragraphs[0].style.name == "Heading 1"
    assert [c.text for c in out.tables[0].rows[0].cells] == ["CELLULE GAUCHE", "CELLULE DROITE"]


def test_formats(tmp_path, gemma):
    import docx as pydocx
    from iak import bookformats
    txt = tmp_path / "conte.txt"
    txt.write_text("Le relieur répare l'atlas.\n\nIl le rend au printemps, { sans } frais.\n",
                   encoding="utf-8")
    job = book.prepare(tmp_path / "lib", settings_for(txt, passes=1, format="epub"))
    assert book.run_job(job, gemma.translator(), lambda: False) == "done"
    assert job.output.name == "conte.en.epub"
    with zipfile.ZipFile(job.output) as z:
        assert z.infolist()[0].filename == "mimetype"
        assert z.infolist()[0].compress_type == zipfile.ZIP_STORED
        opf = z.read("OEBPS/content.opf").decode()
        c1 = z.read("OEBPS/c1.xhtml").decode()
        assert "<dc:language>en</dc:language>" in opf and "properties=\"nav\"" in opf
        assert "<p>LE RELIEUR RÉPARE L'ATLAS.</p>" in c1
    out = {}
    for fmt in ("html", "docx", "rtf", "txt", "pdf"):
        out[fmt] = book.render(job, fmt, tmp_path / f"conte.{fmt}")
    assert "<p>IL LE REND AU PRINTEMPS, { SANS } FRAIS.</p>" in out["html"].read_text("utf-8")
    assert [p.text for p in pydocx.Document(str(out["docx"])).paragraphs][:1] == [
        "LE RELIEUR RÉPARE L'ATLAS."]
    rtf = out["rtf"].read_text("ascii")
    assert rtf.startswith("{\\rtf1") and "R\\u201?PARE" in rtf and "\\{ SANS \\}" in rtf
    assert out["txt"].read_text("utf-8").startswith("LE RELIEUR")
    # EPUB, HTML and DOCX books keep their own format
    with pytest.raises(book.BookError):
        book.output_format("epub", "pdf")
    assert book.output_format("epub", "auto") == "epub"
    assert book.formats_for("pdf") == list(bookformats.FORMATS)
    assert book.formats_for("docx") == ["docx"]
    # chapters cut at the headings
    blocks = [("h1", "Un"), ("p", "a"), ("h1", "Deux"), ("p", "b"), ("li", "c"), ("li", "d")]
    bookformats.write_epub(blocks, "Titre", "fr", tmp_path / "t.epub")
    with zipfile.ZipFile(tmp_path / "t.epub") as z:
        assert {"OEBPS/c1.xhtml", "OEBPS/c2.xhtml"} <= set(z.namelist())
        assert "<ul>\n<li>c</li>\n<li>d</li>\n</ul>" in z.read("OEBPS/c2.xhtml").decode()
        assert ">Deux</a>" in z.read("OEBPS/nav.xhtml").decode()


def test_scanned_pdf_without_ocr(tmp_path):
    pymupdf = pytest.importorskip("pymupdf")
    path = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(path)
    with pytest.raises(book.BookError, match="OCR"):
        book.pdf_paragraphs(path)


def test_glossary_and_trial(tmp_path):
    assert book.parse_glossary("# terms\natlas = map-book\nrelieur\tbookbinder\n\nbad") == [
        ("atlas", "map-book"), ("relieur", "bookbinder")]
    g = FakeGemma()
    try:
        src = make_epub(tmp_path / "r.epub", chapters=3)
        job = book.prepare(tmp_path / "lib", settings_for(src, limit=1, passes=2,
                                                          glossary="atlas = map-book"))
        assert book.run_job(job, g.translator(), lambda: False) == "done"
        assert job.output.name == "r.en.trial.epub" and job.data["units"] == 1
        revisions = [p for p in g.prompts if "REVISED" in p]
        assert revisions and "- atlas → map-book" in revisions[0]
    finally:
        g.close()


def test_books_queue_and_restart(tmp_path, gemma):
    src = make_epub(tmp_path / "r.epub")
    lib = tmp_path / "lib"
    books = book.Books(lib, gemma.translator)
    job = books.create(settings_for(src, passes=1))
    assert books.wait_idle() and job.data["status"] == "done"
    with pytest.raises(KeyError):
        books.get("nothing")
    # a job left "running" (IA-K closed) is interrupted at the next start
    job.save(status="running")
    again = book.Books(lib, gemma.translator)
    assert again.get(job.id).data["status"] == "interrupted"
    again.resume(job.id)
    assert again.wait_idle() and again.get(job.id).data["status"] == "done"
    again.delete(job.id)
    assert not job.folder.exists() and again.list() == []


def test_api_books(tmp_path, txt_file):
    g = FakeGemma(delay=0.05)
    host = servers.Host("test", chat="http://127.0.0.1:9/v1", translate=g.url)
    app = api.App(tmp_path / "lib", host, token=TOKEN)
    srv = api.make_server(app)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    port = srv.server_address[1]
    try:
        with Library(tmp_path / "lib") as lib:
            doc = lib.import_file(txt_file).doc_id
        assert call(port, "POST", "/api/books", {"tgt": "en"})[0] == 400
        assert call(port, "POST", "/api/books", {"doc_id": 99, "tgt": "en"})[0] == 404
        assert call(port, "POST", "/api/books", {"path": str(tmp_path / "x.epub"),
                                                 "tgt": "en"})[0] == 404
        s, b, _ = call(port, "POST", "/api/books", {"doc_id": doc, "src": "fr", "tgt": "en",
                                                    "passes": 1})
        assert s == 202 and b["status"] == "queued" and b["kind"] == "txt"
        assert app.books.wait_idle()
        s, info, _ = call(port, "GET", f"/api/books/{b['id']}")
        assert info["status"] == "done" and info["output_ready"] and info["log"]
        out = tmp_path / "copy.txt"
        s, r, _ = call(port, "POST", f"/api/books/{b['id']}/save", {"path": str(out)})
        assert s == 200 and "DEUXIÈME PARAGRAPHE SUR LA RELIURE." in out.read_text("utf-8")
        assert call(port, "POST", f"/api/books/{b['id']}/save",
                    {"path": str(out), "partial": True})[0] == 409
        # the same book saved as a PDF, and as an EPUB
        assert info["formats"][0] == "pdf" and info["out_format"] == "txt"
        for fmt, magic in (("pdf", b"%PDF"), ("epub", b"PK")):
            dest = tmp_path / f"copy.{fmt}"
            s, _, _ = call(port, "POST", f"/api/books/{b['id']}/save",
                           {"path": str(dest), "format": fmt})
            assert s == 200 and dest.read_bytes().startswith(magic)
        assert call(port, "POST", f"/api/books/{b['id']}/save",
                    {"path": str(dest), "format": "xls"})[0] == 400
        s, r, _ = call(port, "POST", f"/api/books/{b['id']}/import")
        j = wait(port, r["job"])
        assert j["result"]["files"][0]["name"] == "journal.en.txt"
        # an EPUB file from the disk, paused while it runs, then resumed
        epub = make_epub(tmp_path / "long.epub", chapters=8)
        assert call(port, "POST", "/api/books", {"path": str(epub), "tgt": "de",
                                                 "format": "pdf"})[0] == 400
        s, b2, _ = call(port, "POST", "/api/books", {"path": str(epub), "tgt": "de",
                                                     "passes": 2})
        assert s == 202 and b2["src"] == "auto"
        for _ in range(200):
            if app.books.get(b2["id"]).data.get("done_units"):
                break
            time.sleep(0.02)
        call(port, "POST", f"/api/books/{b2['id']}/pause")
        assert app.books.wait_idle()
        _, info, _ = call(port, "GET", f"/api/books/{b2['id']}")
        assert info["status"] == "paused" and info["partial_ready"] and info["src"] == "fr"
        assert call(port, "DELETE", f"/api/books/{b2['id']}")[0] == 200
        _, books, _ = call(port, "GET", "/api/books")
        assert [x["id"] for x in books] == [b["id"]]
        assert call(port, "GET", "/api/books/none")[0] == 404
    finally:
        srv.shutdown()
        srv.server_close()
        g.close()
