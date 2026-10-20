"""Translation of whole books, layout kept: EPUB, PDF, TXT, Markdown, HTML, DOCX.

Ported from truchement's traduire-livre.py (github.com/fhoudebert/transcribe).

1. Cutting. The book is read as SEGMENTS (a paragraph, a heading, a list
   item, a cell...); consecutive segments of a chapter are grouped into
   UNITS of at most max_chars characters, so that the model sees the context
   (dialogue, pronouns, tone) a lone paragraph lacks. The size is capped by
   the server's real context.
2. Passes, for each unit: 1 translation; 2 bilingual revision (original +
   translation -> corrected translation); 3 final proofreading. A revision
   whose output is invalid (tags lost, aberrant length) is rejected: the
   previous pass is kept.
3. Layout (EPUB, HTML). Formatting inside a paragraph (italics, bold, note
   calls, links, page anchors...) is replaced by neutral tags <g1>...</g1>,
   <x2/> the model must keep; the original tags are put back around the
   translated text. When the model damages them: again paragraph by
   paragraph, then plain text as a last resort (counted in the report).
4. Resuming. Every result is appended to cache.jsonl as soon as it is
   obtained; a pause, a shutdown or a crash loses nothing. Each segment's
   final translation is also stored by its own text, so resuming does not
   depend on the cutting, the context or the model.
5. Rebuilding. EPUB -> EPUB (structure, styles, images, table of contents,
   declared language); PDF -> Markdown (headings from font sizes, running
   heads and page numbers removed, paragraphs joined across page breaks,
   hyphenation removed; pages read by OCR in the library are taken from it);
   TXT -> TXT; Markdown -> Markdown; HTML -> HTML; DOCX -> DOCX (formatting
   inside a paragraph simplified).

A book is translated in its own folder, <library>/books/<id>/: job.json
(settings, progress, state), cache.jsonl, book.log, the partial book written
at a pause, and the translated book.
"""
from __future__ import annotations

import copy
import hashlib
import html
import html.entities
import json
import os
import posixpath
import re
import shutil
import threading
import time
import traceback
import zipfile
from collections import Counter
from pathlib import Path
from urllib.parse import unquote

from . import translate
from .servers import ServerError

DEFAULT_MAX_CHARS = 2500
MAX_SEGS_PER_UNIT = 30
FORMATS = {".epub": "epub", ".pdf": "pdf", ".txt": "txt", ".md": "md", ".markdown": "md",
           ".html": "html", ".htm": "html", ".xhtml": "html", ".docx": "docx"}
OUTPUT_EXT = {"epub": ".epub", "pdf": ".md", "units": ".md", "txt": ".txt", "md": ".md",
              "html": ".html", "docx": ".docx"}


class BookError(RuntimeError):
    pass


def lang_name(code: str) -> str:
    return translate.ENGLISH_NAMES.get(code, code)


def _need_bs4():
    try:
        import bs4  # noqa: F401
        return bs4
    except ImportError as exc:
        raise BookError("the beautifulsoup4 package is missing: run setup_python again "
                        "(it installs it)") from exc


class Stats:
    def __init__(self) -> None:
        self.calls = 0
        self.cached = 0
        self.group_fallbacks = 0
        self.format_simplified = 0
        self.revisions_rejected = 0
        self.split_long = 0
        self.suspicious = 0
        self.untranslated = 0
        self.tags_rescued = 0

    def as_dict(self) -> dict:
        return dict(self.__dict__)


# ---------------------------------------------------------------- resume cache

class Cache:
    """One result per line, {"k": key, "p": pass, "v": text}. Appending keeps a
    constant cost on a big book; a last line cut by a power loss is ignored."""

    def __init__(self, path: Path, fresh: bool = False) -> None:
        self.path = path
        self.data: dict[tuple[str, int], str] = {}
        if fresh and path.exists():
            path.unlink()
        if path.exists():
            with open(path, encoding="utf-8", errors="replace") as fh:
                for line in fh:
                    try:
                        rec = json.loads(line)
                        self.data[(rec["k"], int(rec["p"]))] = rec["v"]
                    except (ValueError, KeyError, TypeError):
                        continue
        path.parent.mkdir(parents=True, exist_ok=True)
        self._fh = open(path, "a", encoding="utf-8")
        self._last_sync = time.monotonic()

    def get(self, key: str, p: int) -> str | None:
        return self.data.get((key, p))

    def put(self, key: str, p: int, value: str) -> None:
        self.data[(key, p)] = value
        self._fh.write(json.dumps({"k": key, "p": p, "v": value}, ensure_ascii=False) + "\n")
        self._fh.flush()
        # without fsync a power loss or an unplugged key can lose the last minutes
        if time.monotonic() - self._last_sync > 20:
            self.sync()

    def sync(self) -> None:
        try:
            self._fh.flush()
            os.fsync(self._fh.fileno())
        except (OSError, ValueError):
            pass
        self._last_sync = time.monotonic()

    def close(self) -> None:
        self.sync()
        try:
            self._fh.close()
        except OSError:
            pass


# ---------------------------------------------------------------- model calls

class BookClient:
    """Two ways to call: "chat" through /v1/chat/completions, the instruction
    and the text in one user message; "translategemma" through /completion
    with TranslateGemma's turn written by hand (its Jinja template wants
    language codes the chat API does not pass). Revision passes use the same
    Gemma turn with our own instruction."""

    retry_delay = 5.0

    def __init__(self, tr: translate.Translator, stats: Stats, log) -> None:
        self.client = tr.client
        self.fmt = tr.format
        self.n_ctx = tr.n_ctx or 8192
        self.model = tr.model
        self.stats = stats
        self.log = log

    def _post(self, url: str, payload: dict) -> dict:
        last = None
        for attempt in range(1, 4):
            try:
                return self.client._call(url, payload)
            except ServerError as exc:
                last = exc
                self.log(f"  ! call to the server failed ({attempt}/3): {exc}")
                if attempt < 3:
                    time.sleep(self.retry_delay * attempt)
        raise ServerError(f"the server no longer answers: {last}")

    def max_tokens_for(self, prompt: str, text: str) -> int:
        want = int(len(text) * 0.7) + 256
        room = self.n_ctx - int(len(prompt) * 0.45) - 64
        return max(256, min(want, room))

    def chat(self, prompt: str, text: str, temperature: float) -> tuple[str, bool]:
        payload = {"model": "", "messages": [{"role": "user", "content": prompt}],
                   "temperature": temperature, "top_p": 0.95,
                   "max_tokens": self.max_tokens_for(prompt, text), "stream": False,
                   "cache_prompt": True, "chat_template_kwargs": {"enable_thinking": False}}
        data = self._post(self.client.v1 + "/chat/completions", payload)
        self.stats.calls += 1
        try:
            choice = data["choices"][0]
            out = choice["message"]["content"] or ""
        except (KeyError, IndexError, TypeError) as exc:
            raise ServerError(f"unexpected answer: {str(data)[:200]}") from exc
        return out, choice.get("finish_reason") == "length"

    def gemma_raw(self, user_text: str, text: str, temperature: float) -> tuple[str, bool]:
        prompt = "<start_of_turn>user\n" + user_text + "<end_of_turn>\n<start_of_turn>model\n"
        payload = {"prompt": prompt, "n_predict": self.max_tokens_for(prompt, text),
                   "temperature": temperature, "top_p": 0.95,
                   "stop": ["<end_of_turn>", "<start_of_turn>"], "cache_prompt": True}
        data = self._post(self.client.root + "/completion", payload)
        self.stats.calls += 1
        truncated = data.get("stop_type") == "limit" or bool(data.get("stopped_limit"))
        return data.get("content", ""), truncated


# ---------------------------------------------------------------- instructions

TAG_RULE = (
    "The text contains markup tags such as <seg n=\"1\">…</seg>, <g1>…</g1> "
    "or <x1/>. Keep every tag exactly as written: one <seg> per original "
    "<seg> with the same number, and each <gN>…</gN> pair around the words "
    "that translate the words it surrounded. Never add, drop or renumber tags.")


def has_markup(text: str) -> bool:
    return bool(re.search(r"<(?:seg n=|g\d+>|x\d+/>)", text))


def parse_glossary(text: str) -> list[tuple[str, str]]:
    """Lines "term = translation" (or tab, →, ;); # starts a comment."""
    pairs = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        for sep in ("=", "\t", "→", ";"):
            if sep in line:
                a, b = line.split(sep, 1)
                if a.strip() and b.strip():
                    pairs.append((a.strip(), b.strip()))
                break
    return pairs


def glossary_block(glossary, text: str, tgt: str) -> str:
    lower = text.lower()
    hits = [(a, b) for a, b in glossary if a.lower() in lower]
    if not hits:
        return ""
    lines = "\n".join(f"- {a} → {b}" for a, b in hits)
    return f"Use these {lang_name(tgt)} terms (glossary):\n{lines}\n\n"


def prompt_translate(src: str, tgt: str, text: str, glossary) -> str:
    S, T = lang_name(src), lang_name(tgt)
    rules = [
        f"Translate the following {S} text into {T}. It is an excerpt from a book.",
        f"Translate faithfully and completely into natural, fluent {T}: keep the "
        "meaning, tone, register, dialogue, proper nouns and numbers, and the "
        "paragraph structure.",
        "Never summarize, never omit anything, never add explanations, notes or comments.",
    ]
    if has_markup(text):
        rules.append(TAG_RULE)
    return ("\n".join(rules) + "\n\n" + glossary_block(glossary, text, tgt)
            + f"Output only the {T} translation.\n\n{S} TEXT:\n{text}")


def prompt_translategemma(src: str, tgt: str, text: str) -> str:
    """TranslateGemma's training prompt, word for word (without its turn)."""
    S, T = lang_name(src), lang_name(tgt)
    return (f"You are a professional {S} ({src}) to {T} ({tgt}) translator. Your goal is "
            f"to accurately convey the meaning and nuances of the original {S} text while "
            f"adhering to {T} grammar, vocabulary, and cultural sensitivities.\n"
            f"Produce only the {T} translation, without any additional explanations or "
            f"commentary. Please translate the following {S} text into {T}:\n\n\n{text}")


def prompt_revise(src: str, tgt: str, original: str, translation: str, glossary,
                  final: bool) -> str:
    S, T = lang_name(src), lang_name(tgt)
    if final:
        head = (f"You are proofreading, for the last time, the {T} translation of a {S} "
                f"book. Read it against the original and improve fluency, idiomatic {T} "
                f"and consistency (terms, names, forms of address), while staying strictly "
                f"faithful to the original. Fix any remaining error. Do not rewrite what "
                f"is already good.")
    else:
        head = (f"You are a meticulous bilingual editor revising the {T} translation of a "
                f"{S} book. Compare the translation with the original and correct: "
                f"mistranslations, omissions, additions, false friends, grammar and "
                f"spelling mistakes, and literal or unnatural phrasing. Keep the author's "
                f"style and register. Do not rewrite what is already correct.")
    parts = [head, "Never summarize and never add comments."]
    if has_markup(original):
        parts.append(TAG_RULE)
    return ("\n".join(parts) + "\n\n" + glossary_block(glossary, original, tgt)
            + f"Output only the revised {T} text.\n\n"
            + f"{S} ORIGINAL:\n{original}\n\n{T} TRANSLATION:\n{translation}\n\n"
            + f"REVISED {T} TRANSLATION:")


_PREAMBLE = re.compile(
    r"^(here(?:'s| is)|voici|translation|traduction|revised|révis|corrected|"
    r"übersetzung|traducción)\b[^\n]{0,80}:\s*\n", re.IGNORECASE)


def clean_output(out: str) -> str:
    s = out.replace("<end_of_turn>", "").strip()
    fence = re.fullmatch(r"```[a-zA-Z]*\n(.*?)\n?```", s, re.DOTALL)
    if fence:
        s = fence.group(1).strip()
    s = _PREAMBLE.sub("", s, count=1).strip()
    return re.sub(r"^REVISED [A-Z][a-zA-Z ]* TRANSLATION:\s*", "", s).strip()


def length_ok(source: str, out: str) -> bool:
    """Guard against truncated, summarised or looping outputs."""
    n = len(re.sub(r"<[^>]+>", "", source))
    m = len(re.sub(r"<[^>]+>", "", out))
    if n < 80:
        return 0 < m < max(400, n * 6)
    return 0.35 * n <= m <= 3.0 * n


# ---------------------------------------------------------------- passes

class Engine:
    def __init__(self, client: BookClient, cache: Cache, src: str, tgt: str, passes: int,
                 glossary, stats: Stats, log) -> None:
        self.client, self.cache = client, cache
        self.src, self.tgt, self.passes = src, tgt, passes
        self.glossary, self.stats, self.log = glossary, stats, log

    def key(self, text: str) -> str:
        h = hashlib.sha1()
        h.update(f"{self.client.fmt}|{self.src}|{self.tgt}|".encode())
        h.update(text.encode("utf-8"))
        return h.hexdigest()

    def _call_translate(self, text: str, temperature: float, generic: bool = False):
        c = self.client
        if c.fmt == "translategemma" and generic:
            # explicit instruction (tag rule included) in a Gemma turn: the native
            # one says nothing of tags, which the model then readily drops
            return c.gemma_raw(prompt_translate(self.src, self.tgt, text, self.glossary),
                               text, temperature)
        if c.fmt == "translategemma":
            return c.gemma_raw(prompt_translategemma(self.src, self.tgt, text), text,
                               temperature)
        return c.chat(prompt_translate(self.src, self.tgt, text, self.glossary), text,
                      temperature)

    def untranslated(self, source: str, out: str) -> bool:
        """Did the model copy the original instead of translating it?"""
        if self.src == self.tgt:
            return False

        def norm(t):
            return re.sub(r"\W+", " ", re.sub(r"<[^>]+>", " ", t)).strip().lower()
        a, b = norm(source), norm(out)
        return len(a) > 20 and a == b

    def _call_revise(self, text: str, current: str, final: bool):
        prompt = prompt_revise(self.src, self.tgt, text, current, self.glossary, final)
        if self.client.fmt == "translategemma":
            return self.client.gemma_raw(prompt, text, 0.1)
        return self.client.chat(prompt, text, 0.1)

    def run(self, text: str, valid, generic: bool = False) -> str | None:
        """Translate text with every pass. valid(output) checks the structure
        (tags). None when pass 1 is unusable (truncated or invalid): the caller
        cuts smaller."""
        k = self.key(("generic|" if generic else "") + text)
        cur = self.cache.get(k, 1)
        if cur is not None:
            self.stats.cached += 1
        else:
            for temperature in (0.1, 0.35):
                out, truncated = self._call_translate(text, temperature, generic)
                out = clean_output(out)
                if truncated or not out or not valid(out):
                    return None
                cur = out
                if length_ok(text, out) and not self.untranslated(text, out):
                    break          # otherwise a second try, a little warmer
            if self.untranslated(text, cur):
                self.stats.untranslated += 1
                flat = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text)).strip()[:60]
                self.log(f"  ! block returned UNTRANSLATED by the model: \"{flat}…\"")
            elif not length_ok(text, cur):
                self.stats.suspicious += 1
                self.log(f"  ! suspicious translation length ({len(text)} -> {len(cur)} "
                         f"chars): \"{text[:60]}…\"")
            self.cache.put(k, 1, cur)

        for p in range(2, self.passes + 1):
            rev = self.cache.get(k, p)
            if rev is None:
                out, truncated = self._call_revise(text, cur, final=(p == 3))
                out = clean_output(out)
                if truncated or not out or not valid(out) or not length_ok(text, out):
                    self.stats.revisions_rejected += 1
                    out = cur          # revision rejected: keep the previous pass
                self.cache.put(k, p, out)
                rev = out
            else:
                self.stats.cached += 1
            cur = rev
        return cur


# ---------------------------------------------------------------- segments

_TOKEN = re.compile(r"<\s*(/?)\s*([gx])\s*(\d+)\s*(/?)\s*>", re.IGNORECASE)
_SEG = re.compile(r'<seg\s+n\s*=\s*["\']?(\d+)["\']?\s*>(.*?)</seg\s*>',
                  re.DOTALL | re.IGNORECASE)
_LETTER = re.compile(r"[^\W\d_]")


def esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def unesc(text: str) -> str:
    return html.unescape(text)


def split_sentences(text: str, limit: int) -> list[str]:
    """Pieces of at most limit characters, cut at the ends of sentences."""
    chunks, cur = [], ""
    for s in re.split(r"(?<=[.!?…;»\"”])\s+", text):
        while len(s) > limit:                         # an enormous sentence
            cut = s.rfind(" ", 0, limit)
            cut = cut if cut > limit // 2 else limit
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(s[:cut].strip())
            s = s[cut:].strip()
        if cur and len(cur) + 1 + len(s) > limit:
            chunks.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        chunks.append(cur)
    return [c for c in chunks if c]


class Segment:
    """A block of text to translate (paragraph, heading...)."""

    def __init__(self, enc: str, plain: str, apply_fn) -> None:
        self.enc = enc            # text with neutral tags, escaped
        self.plain = plain        # plain text
        self._apply = apply_fn    # (encoded translation | None, plain text | None)

    @property
    def has_tags(self) -> bool:
        return bool(_TOKEN.search(self.enc))

    def inline_ok(self, out: str) -> bool:
        return self._apply(out, None, dry_run=True)

    def apply(self, out: str) -> None:
        self._apply(out, None)

    def apply_plain(self, text: str) -> None:
        self._apply(None, text)


def plain_segment(text: str, setter) -> Segment:
    """A segment without inline formatting; setter(translated text)."""
    def apply(out, plain_text, dry_run=False):
        if dry_run:
            return not _TOKEN.search(out)
        setter(unesc(out) if out is not None else plain_text)
        return True
    return Segment(esc(text), text, apply)


def parse_group(out: str, n: int) -> list[str] | None:
    found = {}
    for m in _SEG.finditer(out):
        i = int(m.group(1))
        if i in found:
            return None
        found[i] = m.group(2).strip()
    if sorted(found) != list(range(1, n + 1)) or any(not v for v in found.values()):
        return None
    return [found[i] for i in range(1, n + 1)]


class Translator:
    # after these CONSECUTIVE failures the model is judged unable to keep the
    # tags: they are no longer sent (each failed attempt costs a full call)
    GROUP_FAIL_LIMIT = 3
    INLINE_FAIL_LIMIT = 6

    def __init__(self, engine: Engine, max_chars: int) -> None:
        self.engine = engine
        self.stats, self.log = engine.stats, engine.log
        self.max_chars = max_chars
        self.grouping = True
        self.inline = True
        self._group_fails = 0
        self._inline_fails = 0

    # The pass cache is keyed by the text of the UNIT sent to the model: when
    # the cutting changes between sessions (block size, context, model), those
    # entries no longer match. Each segment's final translation is therefore
    # also kept, keyed by its own text.
    def _final_key(self, seg: Segment) -> str:
        e = self.engine
        return hashlib.sha1(f"final|{e.src}|{e.tgt}|{seg.enc}".encode("utf-8")).hexdigest()

    def _store_final(self, seg: Segment, value: str, encoded: bool) -> None:
        self.engine.cache.put(self._final_key(seg), self.engine.passes,
                              ("E:" if encoded else "P:") + value)

    def final_of(self, seg: Segment) -> str | None:
        k = self._final_key(seg)
        for p in range(3, self.engine.passes - 1, -1):     # best first
            v = self.engine.cache.get(k, p)
            if v is not None:
                return v
        return None

    def _apply_final(self, seg: Segment, v: str) -> bool:
        kind, value = v[:2], v[2:]
        if kind == "E:" and seg.inline_ok(value):
            seg.apply(value)
            return True
        if kind == "P:":
            seg.apply_plain(value)
            return True
        return False

    def _done(self, seg: Segment, res: str | None, plain: str | None = None) -> None:
        if res is not None:
            seg.apply(res)
            self._store_final(seg, res, True)
        else:
            seg.apply_plain(plain)
            self._store_final(seg, plain, False)

    def _run_tagged(self, text: str, valid) -> str | None:
        """Tagged text: native instruction, then the explicit one (TranslateGemma)."""
        res = self.engine.run(text, valid)
        if res is None and self.engine.client.fmt == "translategemma":
            res = self.engine.run(text, valid, generic=True)
            if res is not None:
                self.stats.tags_rescued += 1
        return res

    def _inline_result(self, ok: bool) -> None:
        self._inline_fails = 0 if ok else self._inline_fails + 1
        if self.inline and self._inline_fails >= self.INLINE_FAIL_LIMIT:
            self.inline = False
            self.log("  ! the model does not keep the formatting tags: plain text from "
                     "now on (italics, notes... simplified)")

    def plain_text(self, text: str, depth: int = 0) -> str:
        """Plain text, cut smaller as long as the model truncates its output."""
        limit = self.max_chars if depth == 0 else max(200, self.max_chars >> depth)
        if len(text) > limit:
            self.stats.split_long += 1 if depth == 0 else 0
            return " ".join(self.plain_text(c, depth + 1) for c in split_sentences(text, limit))
        res = self.engine.run(esc(text), lambda o: True)
        if res is not None:
            return unesc(res)
        if depth < 4 and len(text) > 200:
            half = split_sentences(text, max(200, len(text) // 2 + 1))
            if len(half) > 1:
                return " ".join(self.plain_text(c, depth + 1) for c in half)
        raise ServerError("the model truncates its output even on a small block "
                          "(context of the server too small?)")

    def single(self, seg: Segment) -> None:
        final = self.final_of(seg)
        if final is not None and self._apply_final(seg, final):
            self.stats.cached += 1
            return
        if len(seg.enc) > self.max_chars * 1.2 or (seg.has_tags and not self.inline):
            if seg.has_tags:
                self.stats.format_simplified += 1
            self._done(seg, None, self.plain_text(seg.plain))
            return
        if seg.has_tags:
            res = self._run_tagged(seg.enc, seg.inline_ok)
            self._inline_result(res is not None)
        else:
            res = self.engine.run(seg.enc, seg.inline_ok)
        if res is not None:
            self._done(seg, res)
            return
        if seg.has_tags:
            self.stats.format_simplified += 1
        self._done(seg, None, self.plain_text(seg.plain))

    def unit(self, segs: list[Segment]) -> None:
        # segments already translated (earlier session, other cutting) are
        # applied as they are; only the others go to the model
        pending = []
        for seg in segs:
            final = self.final_of(seg)
            if final is not None and self._apply_final(seg, final):
                self.stats.cached += 1
            else:
                pending.append(seg)
        if not pending:
            return
        segs = pending
        if len(segs) == 1 or not self.grouping:
            for s in segs:
                self.single(s)
            return
        text = "\n".join(f'<seg n="{i}">{s.enc}</seg>' for i, s in enumerate(segs, start=1))
        res = self._run_tagged(text, lambda o: parse_group(o, len(segs)) is not None)
        if res is None:
            self.stats.group_fallbacks += 1
            self._group_fails += 1
            if self._group_fails >= self.GROUP_FAIL_LIMIT:
                self.grouping = False
                self.log("  ! the model does not keep grouped paragraphs: paragraph by "
                         "paragraph from now on")
            for s in segs:
                self.single(s)
            return
        self._group_fails = 0
        for s, part in zip(segs, parse_group(res, len(segs))):
            if s.inline_ok(part):
                self._done(s, part)
            else:
                self.single(s)


def group_units(segs: list[Segment], max_chars: int) -> list[list[Segment]]:
    units, cur, size = [], [], 0
    for s in segs:
        n = len(s.enc) + 20
        if cur and (size + n > max_chars or len(cur) >= MAX_SEGS_PER_UNIT):
            units.append(cur)
            cur, size = [], 0
        cur.append(s)
        size += n
    if cur:
        units.append(cur)
    return units


# ---------------------------------------------------------------- EPUB, HTML

def local(name: str | None) -> str:
    return (name or "").split(":")[-1].lower()


BLOCK = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "dt", "dd", "td", "th", "caption",
         "figcaption", "blockquote", "div", "title"}
STRUCT = BLOCK | {"ul", "ol", "dl", "table", "thead", "tbody", "tfoot", "tr", "section",
                  "article", "aside", "nav", "header", "footer", "figure", "body", "hr",
                  "pre", "main", "hgroup", "details", "summary", "address", "form",
                  "fieldset"}
SKIP = {"script", "style", "pre", "code", "math", "svg", "kbd", "samp", "var"}
VOID = {"br", "img", "wbr", "hr", "image", "input", "col", "source"}


def _fix_entities(raw: bytes) -> bytes:
    """&nbsp;, &mdash;... come from the XHTML DTD, which lxml does not load: the
    XML parser would drop them. They become numeric references."""
    text = raw.decode("utf-8", errors="replace")

    def rep(m):
        name = m.group(1)
        if name in ("amp", "lt", "gt", "quot", "apos"):
            return m.group(0)
        cp = html.entities.name2codepoint.get(name)
        return f"&#{cp};" if cp else m.group(0)
    return re.sub(r"&([A-Za-z][A-Za-z0-9]*);", rep, text).encode("utf-8")


def parse_xhtml(raw: bytes):
    from bs4 import BeautifulSoup
    try:
        soup = BeautifulSoup(_fix_entities(raw), "xml")
        if soup.find(lambda t: local(t.name) == "body") is not None:
            return soup, "xml"
    except Exception:  # noqa: BLE001 - any parse failure: try as HTML
        pass
    return BeautifulSoup(raw, "html.parser"), "html"


def make_segment(el) -> Segment:
    """Encode the content of el with numbered neutral tags."""
    from bs4 import NavigableString, Tag
    from bs4.element import CData, Comment, Declaration, Doctype, ProcessingInstruction

    slots: dict[int, tuple[str, object]] = {}
    parts: list[str] = []
    counter = [0]

    def walk(node) -> None:
        for child in node.children:
            if isinstance(child, (Comment, CData, ProcessingInstruction, Declaration, Doctype)):
                continue
            if isinstance(child, NavigableString):
                parts.append(esc(str(child)))
            elif isinstance(child, Tag):
                counter[0] += 1
                n = counter[0]
                name = local(child.name)
                if name in VOID or name in SKIP or not child.get_text().strip():
                    slots[n] = ("x", copy.copy(child))
                    parts.append(f"<x{n}/>")
                else:
                    shell = copy.copy(child)
                    shell.clear()
                    slots[n] = ("g", shell)
                    parts.append(f"<g{n}>")
                    walk(child)
                    parts.append(f"</g{n}>")

    walk(el)
    # [ \t\r\n] and not \s: no-break spaces (typography) are kept
    enc = re.sub(r"[ \t\r\n]+", " ", "".join(parts)).strip(" ")
    plain = re.sub(r"\s+", " ", el.get_text(" ")).strip()

    def decode(out: str):
        """The list of nodes again; None when the tags are damaged."""
        root: list = []
        stack: list[tuple[int | None, list, object]] = [(None, root, None)]
        used: set[int] = set()
        pos = 0
        for m in _TOKEN.finditer(out):
            text = out[pos:m.start()]
            if text:
                stack[-1][1].append(NavigableString(unesc(text)))
            pos = m.end()
            closing, kind, n = m.group(1), m.group(2).lower(), int(m.group(3))
            if n not in slots or slots[n][0] != kind:
                return None
            if kind == "x":
                if n in used:
                    return None
                used.add(n)
                stack[-1][1].append(copy.copy(slots[n][1]))
            elif closing:
                if stack[-1][0] != n:
                    return None
                _, children, tag = stack.pop()
                for c in children:
                    tag.append(c)
                stack[-1][1].append(tag)
            else:
                if n in used:
                    return None
                used.add(n)
                stack.append((n, [], copy.copy(slots[n][1])))
        tail = out[pos:]
        if tail:
            stack[-1][1].append(NavigableString(unesc(tail)))
        if len(stack) != 1:
            return None
        missing = set(slots) - used
        # a <gN> pair or a meaningful element (line break, image) lost: invalid;
        # only an empty anchor can be put back
        if any(slots[n][0] == "g" or local(slots[n][1].name) in VOID | SKIP for n in missing):
            return None
        # forgotten empty anchors (page marks, note targets) go first rather
        # than being lost: the links to them still work
        return [copy.copy(slots[n][1]) for n in sorted(missing)] + root

    def apply(out, plain_text, dry_run=False):
        if out is not None:
            nodes = decode(out)
            if dry_run:
                return nodes is not None
            if nodes is None:
                raise ValueError("invalid tags")
        else:
            anchors = [copy.copy(t) for k, t in slots.values()
                       if k == "x" and local(t.name) not in VOID | SKIP]
            nodes = anchors + [NavigableString(plain_text)]
        el.clear()
        for node in nodes:
            el.append(node)
        return True

    return Segment(enc, plain, apply)


def html_targets(soup) -> list:
    targets = []
    for el in soup.find_all(True):
        name = local(el.name)
        if name not in BLOCK:
            continue
        if any(local(a.name) in SKIP for a in el.parents):
            continue
        if name == "title" and not any(local(a.name) == "head" for a in el.parents):
            continue
        if el.find(lambda t: local(t.name) in STRUCT) is not None:
            # table of contents entry: <li><a>Title</a><ol>…</ol></li>
            if name == "li":
                for child in el.find_all(recursive=False):
                    if local(child.name) in ("a", "span") and _LETTER.search(child.get_text()):
                        targets.append(child)
            continue
        if _LETTER.search(el.get_text()):
            targets.append(el)
    return targets


def _set_lang(soup, tgt: str) -> None:
    root = soup.find(lambda t: local(t.name) == "html")
    if root is not None:
        for attr in ("lang", "xml:lang"):
            if attr in root.attrs:
                root[attr] = tgt


def epub_layout(zin: zipfile.ZipFile) -> tuple[str | None, list[str], list[str]]:
    """(path of the OPF, documents in reading order, NCX files)."""
    from bs4 import BeautifulSoup
    names = set(zin.namelist())
    opf = None
    try:
        cont = BeautifulSoup(zin.read("META-INF/container.xml"), "xml")
        rf = cont.find(lambda t: local(t.name) == "rootfile")
        if rf and rf.get("full-path") in names:
            opf = rf["full-path"]
    except KeyError:
        pass
    if opf is None:
        opf = next((n for n in zin.namelist() if n.lower().endswith(".opf")), None)
    html_docs = [n for n in zin.namelist() if n.lower().endswith((".xhtml", ".html", ".htm"))]
    if opf is None:
        return None, html_docs, []
    soup = BeautifulSoup(zin.read(opf), "xml")
    base = posixpath.dirname(opf)
    manifest, ncx = {}, []
    for item in soup.find_all(lambda t: local(t.name) == "item"):
        href = item.get("href")
        if not href:
            continue
        path = posixpath.normpath(posixpath.join(base, unquote(href)))
        manifest[item.get("id")] = path
        if item.get("media-type") == "application/x-dtbncx+xml":
            ncx.append(path)
    order = []
    for ref in soup.find_all(lambda t: local(t.name) == "itemref"):
        p = manifest.get(ref.get("idref"))
        if p in names and p not in order:
            order.append(p)
    order += [d for d in html_docs if d not in order]
    return opf, order, [n for n in ncx if n in names]


class EpubBook:
    kind = "epub"

    def __init__(self, path: Path, tgt: str, max_chars: int) -> None:
        _need_bs4()
        from bs4 import BeautifulSoup
        self.tgt = tgt
        try:
            with zipfile.ZipFile(path) as zin:
                self.infos = zin.infolist()
                self.raw = {i.filename: zin.read(i.filename) for i in self.infos}
                self.opf, order, ncx = epub_layout(zin)
        except zipfile.BadZipFile as exc:
            raise BookError(f"not a valid EPUB: {exc}") from exc
        self.soups: dict[str, object] = {}
        self.chapters: list[tuple[str, list[list[Segment]]]] = []
        for name in order:
            soup, _ = parse_xhtml(self.raw[name])
            segs = [make_segment(el) for el in html_targets(soup)]
            self.soups[name] = soup
            if segs:
                self.chapters.append((name, group_units(segs, max_chars)))
        for name in ncx:
            soup = BeautifulSoup(_fix_entities(self.raw[name]), "xml")
            segs = []
            for t in soup.find_all(lambda t: local(t.name) == "text"):
                if _LETTER.search(t.get_text()):
                    segs.append(self._ncx_segment(t))
            self.soups[name] = soup
            if segs:
                self.chapters.append((name, group_units(segs, max_chars)))

    @staticmethod
    def _ncx_segment(t) -> Segment:
        text = re.sub(r"\s+", " ", t.get_text()).strip()

        def setter(value):
            t.string = value
        return plain_segment(text, setter)

    def write(self, out: Path) -> None:
        data = {}
        for name, soup in self.soups.items():
            _set_lang(soup, self.tgt)
            data[name] = str(soup).encode("utf-8")
        if self.opf and self.opf in self.raw:
            text = self.raw[self.opf].decode("utf-8", errors="replace")
            text = re.sub(r"(<dc:language[^>]*>)[^<]*(</dc:language>)",
                          rf"\g<1>{self.tgt}\g<2>", text)
            data[self.opf] = text.encode("utf-8")
        out.parent.mkdir(parents=True, exist_ok=True)
        tmp = out.with_name(out.name + ".part")
        with zipfile.ZipFile(tmp, "w") as zout:
            # "mimetype" first and stored, as the standard requires
            for info in sorted(self.infos, key=lambda i: i.filename != "mimetype"):
                payload = data.get(info.filename, self.raw[info.filename])
                kind = zipfile.ZIP_STORED if info.filename == "mimetype" else zipfile.ZIP_DEFLATED
                zout.writestr(info, payload, compress_type=kind)
        os.replace(tmp, out)


class HtmlBook:
    kind = "html"

    def __init__(self, path: Path, tgt: str, max_chars: int) -> None:
        _need_bs4()
        self.tgt = tgt
        self.soup, self.mode = parse_xhtml(path.read_bytes())
        segs = [make_segment(el) for el in html_targets(self.soup)]
        self.chapters = [(path.name, group_units(segs, max_chars))]

    def write(self, out: Path) -> None:
        _set_lang(self.soup, self.tgt)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(str(self.soup), encoding="utf-8")


# ---------------------------------------------------------------- PDF, text

_END_SENTENCE = re.compile(r"[.!?:;»\"”)\]…]$")


def join_lines(a: str, b: str) -> str:
    """Join two lines, removing a hyphenation."""
    if not a:
        return b
    if a.endswith("-") and b[:1].islower() and len(a) > 1 and a[-2].isalpha():
        return a[:-1] + b
    return a + " " + b


class Para:
    def __init__(self, kind: str, text: str, prefix: str = "") -> None:
        self.kind = kind          # p, h1, h2, code
        self.text = text
        self.prefix = prefix      # Markdown: "# ", "- ", "> "...
        self.out: str | None = None

    def segment(self) -> Segment:
        def setter(value):
            self.out = value
        return plain_segment(self.text, setter)


def _text_blocks(text: str) -> list[str]:
    """Paragraphs of a page's text: blank lines separate them, lines are joined."""
    out = []
    for block in re.split(r"\n\s*\n", text):
        joined = ""
        for line in block.splitlines():
            line = re.sub(r"\s+", " ", line).strip()
            if line:
                joined = join_lines(joined, line)
        if joined:
            out.append(joined)
    return out


def pdf_paragraphs(path: Path | None, page_texts: dict[int, str] | None = None,
                   n_pages: int = 0) -> list[Para]:
    """Paragraphs of a PDF. page_texts {page number (1-based): text} replaces
    the PDF's own text for those pages (OCR or corrected in the library);
    without a file, the book is made of page_texts alone."""
    page_texts = page_texts or {}
    pages: list[list[tuple[str, float | None]]] = []
    sizes: Counter = Counter()
    if path is not None:
        import pymupdf
        with pymupdf.open(path) as doc:
            for i, page in enumerate(doc, 1):
                if i in page_texts:
                    pages.append([(t, None) for t in _text_blocks(page_texts[i])])
                    continue
                blocks = []
                for b in page.get_text("dict").get("blocks", []):
                    if b.get("type") != 0:
                        continue
                    text, weight, nch = "", 0.0, 0
                    for line in b.get("lines", []):
                        lt = re.sub(r"\s+", " ", "".join(s.get("text", "")
                                                         for s in line.get("spans", []))).strip()
                        for s in line.get("spans", []):
                            n = len(s.get("text", "").strip())
                            weight += s.get("size", 0) * n
                            nch += n
                            sizes[round(s.get("size", 0), 1)] += n
                        if lt:
                            text = join_lines(text, lt)
                    if text:
                        blocks.append((text, weight / max(nch, 1)))
                pages.append(blocks)
    else:
        for i in range(1, max(n_pages, max(page_texts, default=0)) + 1):
            pages.append([(t, None) for t in _text_blocks(page_texts.get(i, ""))])

    body = sizes.most_common(1)[0][0] if sizes else 10.0

    def norm(t: str) -> str:
        return re.sub(r"\d+", "#", t.lower()).strip()

    running: set[str] = set()
    if len(pages) >= 4:
        cnt: Counter = Counter()
        for blocks in pages:
            cnt.update({norm(t) for t, _ in blocks[:2] + blocks[-2:] if len(t) < 120})
        running = {k for k, v in cnt.items() if v >= max(3, 0.3 * len(pages))}

    paras: list[Para] = []
    for blocks in pages:
        for text, size in blocks:
            if re.fullmatch(r"[\d\s\-–—.|]+|(page|p\.)\s*\d+", text, re.I):
                continue                                   # page number
            if norm(text) in running:
                continue                                   # running head or foot
            kind = "p"
            if size is not None and size >= body * 1.6 and len(text) < 150:
                kind = "h1"
            elif size is not None and size >= body * 1.2 and len(text) < 200:
                kind = "h2"
            prev = paras[-1] if paras else None
            if (kind == "p" and prev is not None and prev.kind == "p"
                    and not _END_SENTENCE.search(prev.text) and text[:1].islower()):
                prev.text = join_lines(prev.text, text)    # continued after a page break
                continue
            paras.append(Para(kind, text))
    if not paras:
        raise BookError("no text in this PDF (scanned?): import it into the library, "
                        "where OCR reads it, and translate it from there")
    return paras


_MD_PREFIX = re.compile(r"^(\s*(?:#{1,6}|[-*+]|\d{1,3}[.)]|>+)\s+)(.*)$")


def markdown_paragraphs(text: str) -> list[Para]:
    paras: list[Para] = []
    blocks = re.split(r"(```.*?```)", text, flags=re.S)
    for block in blocks:
        if block.startswith("```"):
            paras.append(Para("code", block))
            continue
        for chunk in re.split(r"\n\s*\n", block):
            lines = [l for l in chunk.splitlines() if l.strip()]
            if not lines:
                continue
            prefixed = [_MD_PREFIX.match(l) for l in lines]
            if len(lines) > 1 and all(prefixed):          # a list, quoted lines
                for m in prefixed:
                    paras.append(Para("p", m.group(2).strip(), m.group(1)))
                continue
            m = prefixed[0]
            prefix, first = (m.group(1), m.group(2)) if m else ("", lines[0])
            joined = first.strip()
            for l in lines[1:]:
                joined = join_lines(joined, l.strip())
            paras.append(Para("p", joined, prefix))
    return paras


def _group_markdown_lists(paras: list[Para]) -> list[str]:
    """Output lines: list items stay on consecutive lines."""
    out: list[str] = []
    prev_item = False
    for p in paras:
        text = p.out if p.out is not None else p.text
        item = bool(p.prefix) and not p.prefix.lstrip().startswith("#")
        line = text if p.kind == "code" else p.prefix + text
        if item and prev_item:
            out[-1] += "\n" + line
        else:
            out.append(line)
        prev_item = item
    return out


class FlatBook:
    """PDF -> Markdown, Markdown -> Markdown, TXT -> TXT."""

    def __init__(self, kind: str, paras: list[Para], name: str, max_chars: int) -> None:
        self.kind = kind
        self.paras = paras
        segs = [p.segment() for p in paras if p.kind != "code" and _LETTER.search(p.text)]
        self.chapters = [(name, group_units(segs, max_chars))]

    def write(self, out: Path) -> None:
        if self.kind == "md":
            lines = _group_markdown_lists(self.paras)
        else:
            lines = []
            for p in self.paras:
                text = p.out if p.out is not None else p.text
                if self.kind in ("pdf", "units") and p.kind == "h1":
                    text = "# " + text
                elif self.kind in ("pdf", "units") and p.kind == "h2":
                    text = "## " + text
                lines.append(text)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text("\n\n".join(lines) + "\n", encoding="utf-8")


class DocxBook:
    """DOCX -> DOCX. Each paragraph (tables included) is translated as a whole;
    its first run takes the translation, so formatting inside a paragraph is
    simplified to the first run's."""
    kind = "docx"

    def __init__(self, path: Path, max_chars: int) -> None:
        import docx
        self.doc = docx.Document(str(path))
        segs = []
        for para in self._paragraphs():
            text = para.text.strip()
            if text and _LETTER.search(text):
                segs.append(plain_segment(re.sub(r"\s+", " ", text), self._setter(para)))
        self.chapters = [(path.name, group_units(segs, max_chars))]

    def _paragraphs(self):
        yield from self.doc.paragraphs
        for table in self.doc.tables:
            for row in table.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:            # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    yield from cell.paragraphs

    @staticmethod
    def _setter(para):
        def setter(value):
            runs = [r for r in para.runs if r.text]
            if not runs:
                para.add_run(value)
                return
            runs[0].text = value
            for r in runs[1:]:
                r.text = ""
        return setter

    def write(self, out: Path) -> None:
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))


def open_book(path: Path | None, kind: str, tgt: str, max_chars: int,
              page_texts: dict[int, str] | None = None, n_pages: int = 0, name: str = ""):
    name = name or (path.name if path else "book")
    if kind == "epub":
        return EpubBook(path, tgt, max_chars)
    if kind == "html":
        return HtmlBook(path, tgt, max_chars)
    if kind == "docx":
        return DocxBook(path, max_chars)
    if kind in ("pdf", "units"):
        return FlatBook(kind, pdf_paragraphs(path if kind == "pdf" else None, page_texts,
                                             n_pages), name, max_chars)
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if kind == "md":
        return FlatBook("md", markdown_paragraphs(text), name, max_chars)
    paras = [Para("p", re.sub(r"\s+", " ", p).strip())
             for p in re.split(r"\n\s*\n", text) if p.strip()]
    return FlatBook("txt", paras, name, max_chars)


def book_sample(book, limit: int = 6000) -> str:
    parts, n = [], 0
    for _, units in book.chapters:
        for unit in units:
            for s in unit:
                parts.append(s.plain)
                n += len(s.plain)
                if n > limit:
                    return " ".join(parts)
    return " ".join(parts)


# ---------------------------------------------------------------- a book job

def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def fmt_duration(sec: float) -> str:
    sec = int(max(0, sec))
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h} h {m:02d}"
    return f"{m} min {s:02d}" if m else f"{s} s"


def _slug(name: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").lower()
    return s[:40] or "book"


ACTIVE = ("queued", "running")


class Job:
    """A book's folder: job.json, cache.jsonl, book.log, outputs."""

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        self.id = folder.name
        try:
            self.data = json.loads((folder / "job.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.data = {}
        self._last_save = 0.0
        self._lock = threading.Lock()
        self.echo = None          # also print the log (command line)

    def save(self, force: bool = True, **fields) -> None:
        with self._lock:
            self.data.update(fields)
            now = time.monotonic()
            if not force and now - self._last_save < 10:
                return
            self._last_save = now
            self.data["updated"] = time.time()
            tmp = self.folder / "job.json.tmp"
            try:
                tmp.write_text(json.dumps(self.data, ensure_ascii=False, indent=1),
                               encoding="utf-8")
                os.replace(tmp, self.folder / "job.json")
            except OSError:
                pass

    def log(self, msg: str) -> None:
        if self.echo:
            self.echo(msg)
        try:
            with open(self.folder / "book.log", "a", encoding="utf-8") as fh:
                fh.write(time.strftime("%Y-%m-%d %H:%M:%S  ") + msg + "\n")
        except OSError:
            pass

    def log_tail(self, n: int = 40) -> list[str]:
        try:
            lines = (self.folder / "book.log").read_text(encoding="utf-8",
                                                         errors="replace").splitlines()
        except OSError:
            return []
        return lines[-n:]

    @property
    def output(self) -> Path | None:
        name = self.data.get("output")
        return self.folder / name if name else None

    @property
    def partial(self) -> Path | None:
        name = self.data.get("partial")
        return self.folder / name if name else None

    def public(self) -> dict:
        d = {k: v for k, v in self.data.items() if k not in ("page_texts", "glossary")}
        d["id"] = self.id
        d["glossary_entries"] = len(parse_glossary(self.data.get("glossary", "")))
        d["output_ready"] = bool(self.output and self.output.exists()
                                 and self.data.get("status") == "done")
        d["partial_ready"] = bool(self.partial and self.partial.exists())
        return d


class PauseRequested(Exception):
    pass


def run_job(job: Job, translator: translate.Translator, should_pause) -> str:
    """Translate the book of job; returns the final status (done, paused,
    error). Every step is in job.json and book.log."""
    d = job.data
    stats = Stats()
    log = job.log
    kind = d["kind"]
    source = Path(d["source"]) if d.get("source") else None
    if source is not None and not source.is_file():
        raise BookError(f"the book is no longer there: {source}")
    page_texts = {int(k): v for k, v in (d.get("page_texts") or {}).items()}

    client = BookClient(translator, stats, log)
    # unit size capped by the real context: a revision carries the original and
    # the translation as input AND the revision as output
    cap = max(600, int((client.n_ctx * 0.85 - 400) / 1.3))
    max_chars = max(300, min(int(d.get("max_chars") or DEFAULT_MAX_CHARS), cap))
    book = open_book(source, kind, d["tgt"], max_chars, page_texts, d.get("pages", 0),
                     d.get("name", ""))
    src = d.get("src") or "auto"
    if src == "auto":
        src = translate.detect_language(book_sample(book))
        if not src:
            raise BookError("the language of the book could not be detected: choose it")
        log(f"Detected language: {lang_name(src)}")
    tgt, passes = d["tgt"], int(d.get("passes") or 2)
    glossary = parse_glossary(d.get("glossary", ""))
    log("-" * 60)
    log(f"Book: {d.get('name')}  ·  {lang_name(src)} -> {lang_name(tgt)}  ·  {passes} pass(es)")
    log(f"Server: {translator.client.root}  ·  {translator.model}  ·  context "
        f"{client.n_ctx}  ·  mode {client.fmt}  ·  blocks of {max_chars} chars")
    if glossary:
        log(f"Glossary: {len(glossary)} entries"
            + ("  (TranslateGemma: used by the revision passes)"
               if client.fmt == "translategemma" else ""))

    cache = Cache(job.folder / "cache.jsonl", fresh=bool(d.pop("fresh", False)))
    engine = Engine(client, cache, src, tgt, passes, glossary, stats, log)
    tr = Translator(engine, max_chars)
    plan = [(name, u) for name, units in book.chapters for u in units]
    full_units = len(plan)
    limit = int(d.get("limit") or 0)
    if limit:
        plan = plan[:limit]
    unit_chars = [sum(len(s.plain) for s in u) for _, u in plan]
    total = sum(unit_chars) or 1
    already = sum(1 for _, u in plan if all(tr.final_of(s) is not None for s in u))
    n_chapters = len({name for name, _ in plan})
    log(f"Cut: {n_chapters} chapter(s), {sum(len(u) for _, u in plan)} segments, "
        f"{len(plan)} units, {total} chars" + (f"; {already} units already translated"
                                                if already else ""))
    if limit and limit < full_units:
        log(f"TRIAL: {len(plan)} unit(s) of {full_units}; the rest stays in the original "
            f"language")
    stem = Path(d.get("name") or "book").stem
    trial = ".trial" if limit and limit < full_units else ""
    out_name = f"{stem}.{tgt}{trial}{OUTPUT_EXT[kind]}"
    job.save(src=src, src_detected=d.get("src") in (None, "", "auto"), mode=client.fmt,
             model=translator.model, n_ctx=client.n_ctx, block_chars=max_chars,
             units=len(plan), done_units=already, chars=total,
             done_chars=sum(n for (_, u), n in zip(plan, unit_chars)
                            if all(tr.final_of(s) is not None for s in u)),
             status="running", error="", started=time.time())

    t0 = time.monotonic()
    before = int(d.get("work_seconds") or 0)       # earlier sessions
    done_chars = fresh_chars = 0
    i_done = 0
    current = None
    try:
        for i, ((name, unit), nch) in enumerate(zip(plan, unit_chars), start=1):
            if should_pause():
                raise PauseRequested
            if name != current:
                current = name
                if n_chapters > 1:
                    log(f"[{len({n for n, _ in plan[:i]})}/{n_chapters}] {name}")
            calls_before = stats.calls
            tr.unit(unit)
            i_done = i
            done_chars += nch
            if stats.calls > calls_before:
                fresh_chars += nch
            elapsed = time.monotonic() - t0
            eta = elapsed / fresh_chars * (total - done_chars) if fresh_chars else None
            job.save(force=(i == len(plan)), done_units=i, done_chars=done_chars,
                     pct=int(100 * done_chars / total), elapsed=int(elapsed),
                     work_seconds=before + int(elapsed),
                     eta=int(eta) if eta is not None else None, stats=stats.as_dict())
    except PauseRequested:
        cache.close()
        partial = None
        if i_done:
            try:
                partial = job.folder / f"{stem}.{tgt}.partial{OUTPUT_EXT[kind]}"
                book.write(partial)
            except Exception as exc:  # noqa: BLE001 - the pause must not fail
                log(f"  (partial book not written: {exc})")
                partial = None
        pct = int(100 * done_chars / total)
        job.save(status="paused", done_units=i_done, done_chars=done_chars, pct=pct,
                 eta=None, work_seconds=before + int(time.monotonic() - t0), partial=partial.name if partial else "", stats=stats.as_dict())
        log(f"Paused after {i_done}/{len(plan)} units ({pct} %). Resume at any time, "
            f"even after a restart.")
        return "paused"
    except Exception:
        cache.close()
        job.save(done_units=i_done, done_chars=done_chars, pct=int(100 * done_chars / total),
                 eta=None, stats=stats.as_dict(),
                 work_seconds=before + int(time.monotonic() - t0))
        raise

    cache.close()
    book.write(job.folder / out_name)
    if job.partial and job.partial.exists():
        try:
            job.partial.unlink()
        except OSError:
            pass
    elapsed = time.monotonic() - t0
    job.save(status="done", output=out_name, partial="", done_units=len(plan),
             done_chars=total, pct=100, eta=None, elapsed=int(elapsed),
             work_seconds=before + int(elapsed),
             stats=stats.as_dict(), finished=time.time())
    s = stats
    details = [f"{s.calls} model call(s)", f"{s.cached} taken from the cache"]
    for n, what in ((s.group_fallbacks, "group(s) translated again paragraph by paragraph"),
                    (s.format_simplified, "paragraph(s) with simplified formatting"),
                    (s.revisions_rejected, "revision(s) rejected"),
                    (s.suspicious, "translation(s) of suspicious length"),
                    (s.untranslated, "block(s) returned untranslated"),
                    (s.tags_rescued, "tagged block(s) saved by the explicit instruction")):
        if n:
            details.append(f"{n} {what}")
    log(f"Done in {fmt_duration(elapsed)}"
        + (f" ({fmt_duration(before + elapsed)} in all)" if before else "") + f": {out_name}")
    log("  " + " · ".join(details))
    return "done"


def job_id_for(settings: dict) -> str:
    return f"{_slug(Path(settings['name']).stem)}-{settings['sha256'][:8]}-{settings['tgt']}"


def prepare(library: Path, settings: dict) -> Job:
    """The job folder for settings, created or updated (no thread: the
    command line runs it itself)."""
    folder = Path(library) / "books" / job_id_for(settings)
    folder.mkdir(parents=True, exist_ok=True)
    job = Job(folder)
    job.save(**settings, created=job.data.get("created") or time.time(), status="queued",
             error="")
    return job


class Books:
    """Book jobs of a library: one translated at a time, in its own thread,
    so that imports and questions are not held up for hours."""

    def __init__(self, library: Path, translator_factory) -> None:
        self.folder = Path(library) / "books"
        self.translator_factory = translator_factory
        self._cond = threading.Condition()
        self._queue: list[str] = []
        self._current: str | None = None
        self._pause: set[str] = set()
        self._jobs: dict[str, Job] = {}
        # a book left running or queued (the application was closed, the
        # computer stopped) can be resumed; it is not restarted on its own,
        # the model servers may not be there yet
        if self.folder.is_dir():
            for d in self.folder.iterdir():
                if (d / "job.json").exists():
                    job = Job(d)
                    self._jobs[job.id] = job
                    if job.data.get("status") in ACTIVE:
                        job.save(status="interrupted")
                        job.log("Interrupted (IA-K was closed); resume when you like.")
        threading.Thread(target=self._worker, daemon=True, name="iak-books").start()

    # ---------------------------------------------------------- jobs
    def list(self) -> list[dict]:
        with self._cond:
            jobs = list(self._jobs.values())
        out = [j.public() for j in jobs]
        return sorted(out, key=lambda d: -(d.get("created") or 0))

    def get(self, job_id: str) -> Job:
        job = self._jobs.get(job_id)
        if job is None:
            raise KeyError(f"no book {job_id}")
        return job

    def create(self, settings: dict) -> Job:
        """settings: name, kind, source (path or None), sha256, tgt, src,
        passes, glossary, limit, max_chars, page_texts, pages, doc_id."""
        job_id = job_id_for(settings)
        folder = self.folder / job_id
        folder.mkdir(parents=True, exist_ok=True)
        with self._cond:
            job = self._jobs.get(job_id) or Job(folder)
            if job.data.get("status") in ACTIVE:
                raise BookError("this book is already being translated into this language")
            created = job.data.get("created") or time.time()
            job.save(**settings, created=created, status="queued", error="")
            self._jobs[job_id] = job
            self._queue.append(job_id)
            self._cond.notify_all()
        job.log(f"Queued: {settings['name']} -> {settings['tgt']}")
        return job

    def resume(self, job_id: str, fresh: bool = False) -> Job:
        job = self.get(job_id)
        with self._cond:
            if job.data.get("status") in ACTIVE:
                return job
            self._pause.discard(job_id)
            job.save(status="queued", error="", **({"fresh": True} if fresh else {}))
            self._queue.append(job_id)
            self._cond.notify_all()
        return job

    def pause(self, job_id: str) -> Job:
        job = self.get(job_id)
        with self._cond:
            if job_id in self._queue:
                self._queue.remove(job_id)
                job.save(status="paused")
            elif self._current == job_id:
                self._pause.add(job_id)
                job.save(pausing=True)
        return job

    def delete(self, job_id: str) -> None:
        job = self.get(job_id)
        with self._cond:
            if job_id == self._current:
                raise BookError("pause the translation before deleting it")
            if job_id in self._queue:
                self._queue.remove(job_id)
            del self._jobs[job_id]
        shutil.rmtree(job.folder, ignore_errors=True)

    def busy(self) -> bool:
        with self._cond:
            return self._current is not None or bool(self._queue)

    def wait_idle(self, timeout: float = 30) -> bool:
        """For tests: wait until nothing runs nor waits."""
        end = time.time() + timeout
        while time.time() < end:
            if not self.busy():
                return True
            time.sleep(0.05)
        return False

    # ---------------------------------------------------------- worker
    def _worker(self) -> None:
        while True:
            with self._cond:
                while not self._queue:
                    self._cond.wait()
                job_id = self._queue.pop(0)
                job = self._jobs.get(job_id)
                if job is None:
                    continue
                self._current = job_id
            try:
                tr = self.translator_factory()
                status = run_job(job, tr, lambda: job_id in self._pause)
                job.save(status=status, pausing=False)
            except (BookError, ServerError, translate.TranslateError, OSError, KeyError,
                    ValueError) as exc:
                job.save(status="error", error=str(exc), pausing=False)
                job.log(f"Stopped: {exc}. What is translated is kept: resume to go on.")
            except Exception as exc:  # noqa: BLE001 - report, keep the worker alive
                job.save(status="error", error=f"{type(exc).__name__}: {exc}", pausing=False)
                job.log(f"Stopped: {type(exc).__name__}: {exc}")
                traceback.print_exc()
            finally:
                with self._cond:
                    self._pause.discard(job_id)
                    self._current = None
